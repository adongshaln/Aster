"""Read Android-generated CI fixtures with independent desktop document parsers."""
import argparse
import posixpath
import unicodedata
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

parser = argparse.ArgumentParser()
parser.add_argument("directory", type=Path)
args = parser.parse_args()

def sample(ext):
    return next(args.directory.rglob(f"generated-sample.{ext}"))

with fitz.open(sample("pdf")) as pdf:
    # Android/Skia may map identical CJK glyphs to compatibility code points.
    # Compare compatible text after NFKC; visual layout is checked separately.
    text = unicodedata.normalize("NFKC", "\n".join(page.get_text() for page in pdf))
    assert len(pdf) >= 3
    assert "中文文档验证" in text and "END_OF_DOCUMENT" in text
    assert "Document paragraph 90." in text
    pdf[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(args.directory / "pdf-first-page.png")
    print(f"PDF: {len(pdf)} pages, Chinese text and final paragraph present")

word = Document(sample("docx"))
assert word.paragraphs[0].text == "中文文档验证"
assert word.paragraphs[-1].text == "END_OF_DOCUMENT"
assert word.tables[0].cell(1, 0).text == "木材"
assert word.tables[0].cell(1, 1).text == "12.5"
print(f"DOCX: {len(word.paragraphs)} paragraphs, table and final paragraph present")

book = load_workbook(sample("xlsx"), data_only=False)
assert book.sheetnames == ["预算", "人物"]
assert book["预算"]["B2"].value == 12.5
assert book["预算"]["B4"].value == "=SUM(B2:B3)"
assert book["人物"]["A2"].value == "爱丽丝"
assert not book._external_links
print("XLSX: two sheets, numeric cells and formula preserved, no external links")

slides = Presentation(sample("pptx"))
assert len(slides.slides) == 2
text = "\n".join(shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
assert "故事创作计划" in text and "完成第一章" in text
print("PPTX: two editable slides, titles and body text present")

for ext in ("docx", "xlsx", "pptx"):
    with ZipFile(sample(ext)) as package:
        names = set(package.namelist())
        for path in names:
            if path.endswith(".xml") or path.endswith(".rels"):
                root = ET.fromstring(package.read(path))
            if not path.endswith(".rels"):
                continue
            parent = posixpath.dirname(posixpath.dirname(path))
            for relation in root:
                assert relation.get("TargetMode") != "External"
                target = relation.attrib["Target"]
                resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(parent, target))
                assert resolved in names, (ext, path, target)
    print(f"{ext.upper()}: every XML part and package relationship valid")

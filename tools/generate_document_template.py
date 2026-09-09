"""Regenerate the bundled blank OOXML presentation using python-pptx (development only)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parents[1]
presentation = Presentation()
presentation.slide_width = Inches(13.333333)
presentation.slide_height = Inches(7.5)
presentation.slides.add_slide(presentation.slide_layouts[6])
presentation.save(root / "app/src/main/resources/document-templates/presentation.pptx")
